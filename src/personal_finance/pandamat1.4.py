import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

#so far the best one

# === Load data ===
excel_file = "/home/eikov/fin/analysis/Manually-polished/2025/Galutinis_spendings.xlsx"  # <-- Replace with your actual file
sheet_name = "Totals"

df = pd.read_excel(excel_file, sheet_name=sheet_name)

# === Group by Category and Subcategory ===
grouped = df.groupby(['Category', 'Subcategory'], as_index=False)['SUMA'].sum()

# Remove zero or negative values
grouped = grouped[grouped['SUMA'] > 0]

# Sort by SUMA descending
grouped = grouped.sort_values(by='SUMA', ascending=False)

# === Create pie chart ===
fig, ax = plt.subplots(figsize=(8, 8))

# Pie data
sizes = grouped['SUMA']
labels_raw = grouped[['Category', 'Subcategory']].agg(' – '.join, axis=1)

# Pie chart
patches, texts, autotexts = ax.pie(
    sizes,
    autopct='%1.1f%%',
    startangle=90
)

# Generate custom labels with % + numeric sum + category info
total = sizes.sum()
percentages = (sizes / total * 100).round(1)
sums_formatted = sizes.round(2)

legend_labels = [
    f"{p:.1f}% – {s:.2f} – {label}"
    for p, s, label in zip(percentages, sums_formatted, labels_raw)
]

# Sort legend entries by % descending
sorted_indices = percentages.argsort()[::-1]
sorted_labels = [legend_labels[i] for i in sorted_indices]
sorted_patches = [patches[i] for i in sorted_indices]

# Add legend
ax.legend(
    sorted_patches,
    sorted_labels,
    loc="center left",
    bbox_to_anchor=(1, 0.5),
    fontsize='small'
)

# Save the pie chart
chart_path = "output_chart.png"
plt.tight_layout()
plt.savefig(chart_path)
plt.close()

# === Insert chart into Excel ===
wb = load_workbook(excel_file)

# Create or get the "Analysis" sheet
if "Analysis" in wb.sheetnames:
    ws = wb["Analysis"]
else:
    ws = wb.create_sheet("Analysis")

# Add image
img = ExcelImage(chart_path)
img.anchor = "A1"
ws.add_image(img)

# Save updated Excel file
wb.save(excel_file)

print("Chart with enhanced legend saved and embedded into 'Analysis' sheet.")
