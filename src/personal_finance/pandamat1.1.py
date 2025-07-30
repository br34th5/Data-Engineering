import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage

# File paths
excel_file = '/home/eikov/fin/analysis/Manually-polished/2025/Galutinis_spendings.xlsx'  # Replace with actual file name
chart_image_path = '/home/eikov/fin/analysis/Manually-polished/2025/spending_pie_by_category.png'

# 1. Load data from Totals sheet
df = pd.read_excel(excel_file, sheet_name='Totals')

# 2. Clean and group
df['SUMA'] = pd.to_numeric(df['SUMA'], errors='coerce')
df = df.dropna(subset=['SUMA'])
grouped = df.groupby('Category')['SUMA'].sum().sort_values(ascending=False)
grouped = grouped[grouped > 0]

# 3. Create and save pie chart
plt.figure(figsize=(8, 8))
grouped.plot.pie(autopct='%1.1f%%', startangle=90)
plt.title('Spending Distribution by Category')
plt.ylabel('')
plt.tight_layout()
plt.savefig(chart_image_path)
plt.close()

# 4. Insert image into a new sheet in the Excel file
workbook = load_workbook(excel_file)
if 'Analysis' in workbook.sheetnames:
    del workbook['Analysis']  # remove if it already exists (optional)
sheet = workbook.create_sheet('Analysis')

# Insert image
img = OpenpyxlImage(chart_image_path)
img.anchor = 'A1'  # Position in the sheet
sheet.add_image(img)

# Save workbook
workbook.save(excel_file)
print(f"✅ Pie chart inserted into 'Analysis' sheet of: {excel_file}")
